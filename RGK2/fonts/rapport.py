"""Génération du rapport Excel : référentiel clients + consommation mensuelle."""
from datetime import datetime
from io import BytesIO

import pandas as pd
from openpyxl import Workbook
from openpyxl.chart import BarChart, Reference
from openpyxl.formatting.rule import ColorScaleRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter as L
from openpyxl.worksheet.table import Table, TableStyleInfo

import logic

TEAL, GREY, INK = "0F6B5C", "55625C", "16201C"
F = "Calibri"
HEAD = PatternFill("solid", fgColor=TEAL)
HEAD2 = PatternFill("solid", fgColor=GREY)
TOT = PatternFill("solid", fgColor="D6EAE3")
BAND = PatternFill("solid", fgColor="F2F5F3")
thin = Side(style="thin", color="D0D5D2")
EUR = '#,##0.00 "€";-#,##0.00 "€";"-"'
EUR0 = '#,##0 "€";-#,##0 "€";"-"'
DATE = "dd/mm/yyyy"


def _f(size=11, bold=False, color=INK, italic=False):
    return Font(name=F, size=size, bold=bold, color=color, italic=italic)


def _titre(ws, titre, sous_titre):
    ws.sheet_view.showGridLines = False
    ws["A1"] = titre
    ws["A1"].font = _f(16, True)
    ws.row_dimensions[1].height = 24
    ws["A2"] = sous_titre
    ws["A2"].font = _f(10, color=GREY, italic=True)


def _entetes(ws, row, headers, col=1, fill=HEAD, prev_from=None):
    for j, h in enumerate(headers):
        c = ws.cell(row=row, column=col + j, value=h)
        c.font = _f(11, True, "FFFFFF")
        c.fill = HEAD2 if (prev_from is not None and j >= prev_from) else fill
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    ws.row_dimensions[row].height = 30


def _table(ws, nom, ref):
    t = Table(displayName=nom, ref=ref)
    t.tableStyleInfo = TableStyleInfo(name="TableStyleLight1", showRowStripes=True)
    ws.add_table(t)


def _largeurs(ws, d):
    for k, v in d.items():
        ws.column_dimensions[k].width = v


def generer_rapport(client_ids=None, releves=False):
    """Retourne le contenu (bytes) d'un classeur Excel. client_ids : ensemble d'ID à inclure (None = tous)."""
    an, mr = logic.annee(), logic.mois_ref()
    mois_cols = list(range(1, 13))
    entetes_mois = [logic.MOIS[i - 1] + ("" if i <= mr else " (prév.)") for i in mois_cols]
    lignes = logic.lignes_mensuelles()
    cc = logic.conso_clients(lignes)
    clients = logic.q("SELECT * FROM clients ORDER BY id")
    if client_ids is not None:
        clients = clients[clients.id.isin(client_ids)]
        cc = cc[cc.client_id.isin(client_ids)]
        lignes = lignes[lignes.client_id.isin(client_ids)]
    act = lignes[lignes.statut == "Active"]
    offre = {cid: f"{g.code.iloc[-1]} {g.libelle.iloc[-1]}" for cid, g in act[act.element == "Offre"].groupby("client_id")}
    nopt = set(act[act.element == "Option"].client_id)
    ccx = cc.set_index("client_id")
    filtre_txt = "tous les clients" if client_ids is None else f"{len(clients)} client(s) sélectionné(s)"
    sous = (f"Exercice {an} · réel jusqu'à {logic.MOIS_LONG[mr - 1]} {an}, prévision ensuite · {filtre_txt} · "
            f"généré le {datetime.now():%d/%m/%Y à %H:%M}")

    wb = Workbook()

    # ------------------------------------------------------------ Clients
    ws = wb.active
    ws.title = "Clients"
    _titre(ws, "Référentiel clients", sous)
    hd = ["ID", "Nom du client", "Entité", "Pays", "Responsable", "E-mail", "Statut", "Date d'entrée",
          "Offre actuelle", f"Coût / mois ({logic.MOIS[mr - 1]})", "Conso réelle", f"Total {an}"]
    _entetes(ws, 4, hd)
    for i, r in enumerate(clients.itertuples()):
        row = 5 + i
        vals = [r.id, r.nom, r.entite, r.pays, r.responsable, r.email, r.statut, logic.d(r.date_entree),
                offre.get(r.id, "—") + (" + options" if r.id in nopt else ""),
                float(ccx.mois_ref.get(r.id, 0)), float(ccx.reel.get(r.id, 0)), float(ccx.total.get(r.id, 0))]
        for j, v in enumerate(vals):
            c = ws.cell(row=row, column=1 + j, value=v)
            c.font = _f()
            if j == 7:
                c.number_format = DATE
            if j >= 9:
                c.number_format = EUR if j == 9 else EUR0
    last = 4 + max(1, len(clients))
    _table(ws, "tRapportClients", f"A4:L{last}")
    tr = last + 1
    ws.cell(row=tr, column=1, value="TOTAL").font = _f(11, True)
    for col in (10, 11, 12):
        c = ws.cell(row=tr, column=col, value=f"=SUBTOTAL(109,{L(col)}5:{L(col)}{last})")
        c.number_format = EUR if col == 10 else EUR0
        c.font = _f(11, True)
    for col in range(1, 13):
        ws.cell(row=tr, column=col).fill = TOT
        ws.cell(row=tr, column=col).border = Border(top=Side(style="medium", color=TEAL))
    _largeurs(ws, {"A": 9, "B": 30, "C": 16, "D": 12, "E": 16, "F": 34, "G": 10, "H": 13, "I": 32, "J": 16, "K": 15, "L": 15})
    ws.freeze_panes = "C5"

    # ------------------------------------------------------------ Consommation mensuelle
    ws = wb.create_sheet("Consommation")
    _titre(ws, "Consommation mensuelle par client", sous + " · les colonnes « prév. » sont des prévisions")
    hd = ["ID", "Client", "Entité", "Pays"] + entetes_mois + [f"Total {an}"]
    _entetes(ws, 4, hd, prev_from=4 + mr)
    ws.cell(row=4, column=17).fill = HEAD
    conso = cc[cc.total > 0].copy()
    ordre = {p: k for k, p in enumerate(logic.liste("pays"))}
    conso = conso.sort_values(["pays", "client"], key=lambda s: s.map(ordre) if s.name == "pays" else s)
    for i, (_, r) in enumerate(conso.iterrows()):
        row = 5 + i
        vals = [r.client_id, r.client, r.entite, r.pays] + [float(r[m]) for m in mois_cols]
        for j, v in enumerate(vals):
            c = ws.cell(row=row, column=1 + j, value=v)
            c.font = _f(italic=(j >= 4 + mr and j < 16), color=("6B7570" if (j >= 4 + mr and j < 16) else INK))
            if j >= 4:
                c.number_format = EUR0
        c = ws.cell(row=row, column=17, value=f"=SUM(E{row}:P{row})")
        c.number_format = EUR0
        c.font = _f(11, True)
    last = 4 + max(1, len(conso))
    _table(ws, "tRapportConso", f"A4:Q{last}")
    tr = last + 1
    ws.cell(row=tr, column=1, value="TOTAL").font = _f(11, True)
    for col in range(1, 18):
        c = ws.cell(row=tr, column=col)
        c.fill = TOT
        c.border = Border(top=Side(style="medium", color=TEAL))
        if col >= 5:
            c.value = f"=SUBTOTAL(109,{L(col)}5:{L(col)}{last})"
            c.number_format = EUR0
            c.font = _f(11, True)
    if len(conso):
        ws.conditional_formatting.add(f"E5:P{last}", ColorScaleRule(start_type="num", start_value=0, start_color="FFFFFF",
                                                                      end_type="max", end_color="9CC7BB"))
    _largeurs(ws, {"A": 9, "B": 28, "C": 15, "D": 11, "Q": 14})
    for k in range(12):
        ws.column_dimensions[L(5 + k)].width = 11
    ws.freeze_panes = "E5"
    conso_last = last

    # ------------------------------------------------------------ Synthèse
    ws = wb.create_sheet("Synthèse", 0)
    _titre(ws, f"IA Platform — rapport de consommation {an}", sous)
    kpis = [("Consommation réelle", f"=SUM(Consommation!E5:{L(4 + mr)}{conso_last})"),
            ("Prévision fin d'année", f"=SUM(Consommation!Q5:Q{conso_last})"),
            (f"Consommation de {logic.MOIS_LONG[mr - 1]}", f"=SUM(Consommation!{L(4 + mr)}5:{L(4 + mr)}{conso_last})"),
            ("Clients consommateurs", f"=COUNTIF(Consommation!{L(4 + mr)}5:{L(4 + mr)}{conso_last},\">0\")")]
    for k, (lab, f) in enumerate(kpis):
        col = 1 + k * 3
        ws.cell(row=4, column=col, value=lab).font = _f(10, color=GREY)
        c = ws.cell(row=5, column=col, value=f)
        c.font = _f(18, True, TEAL)
        c.number_format = "0" if "Clients" in lab else EUR0
        ws.merge_cells(start_row=5, start_column=col, end_row=5, end_column=col + 1)
    ws.row_dimensions[5].height = 28
    _entetes(ws, 7, ["Mois", "Montant", "Période"])
    for m in mois_cols:
        row = 7 + m
        ws.cell(row=row, column=1, value=logic.MOIS_LONG[m - 1].capitalize()).font = _f()
        c = ws.cell(row=row, column=2, value=f"=Consommation!{L(4 + m)}{conso_last + 1}")
        c.number_format = EUR0
        ws.cell(row=row, column=3, value="Réel" if m <= mr else "Prévision").font = _f(color=INK if m <= mr else "8A8F8C", italic=m > mr)
    ws.cell(row=20, column=1, value="Total").font = _f(11, True)
    c = ws.cell(row=20, column=2, value="=SUM(B8:B19)")
    c.number_format = EUR0
    c.font = _f(11, True)
    for col in (1, 2, 3):
        ws.cell(row=20, column=col).fill = TOT
    ch = BarChart()
    ch.type = "col"
    ch.title = f"Consommation mensuelle {an}"
    ch.add_data(Reference(ws, min_col=2, min_row=7, max_row=19), titles_from_data=True)
    ch.set_categories(Reference(ws, min_col=1, min_row=8, max_row=19))
    ch.series[0].graphicalProperties.solidFill = TEAL
    ch.legend = None
    ch.y_axis.numFmt = '#,##0 "€"'
    ch.y_axis.delete = False
    ch.x_axis.delete = False
    ch.width, ch.height = 18, 7.5
    ws.add_chart(ch, "E7")

    def bloc(start_row, titre, col_src, valeurs):
        ws.cell(row=start_row, column=1, value=titre).font = _f(12, True, TEAL)
        _entetes(ws, start_row + 1, [titre.split()[-1].capitalize(), "Clients", "Conso réelle", f"Total {an}"], fill=HEAD2)
        for k, v in enumerate(valeurs):
            row = start_row + 2 + k
            ws.cell(row=row, column=1, value=v).font = _f()
            ws.cell(row=row, column=2, value=f'=COUNTIFS(Consommation!{col_src}5:{col_src}{conso_last},A{row})')
            c = ws.cell(row=row, column=3, value=f"=SUMPRODUCT((Consommation!{col_src}5:{col_src}{conso_last}=A{row})*Consommation!E5:{L(4 + mr)}{conso_last})")
            c.number_format = EUR0
            c = ws.cell(row=row, column=4, value=f'=SUMIFS(Consommation!Q5:Q{conso_last},Consommation!{col_src}5:{col_src}{conso_last},A{row})')
            c.number_format = EUR0
        return start_row + 2 + len(valeurs)

    nxt = bloc(23, "Par pays", "D", sorted(conso.pays.unique(), key=lambda p: ordre.get(p, 99)))
    bloc(nxt + 2, "Par entité", "C", sorted(conso.entite.unique()))
    _largeurs(ws, {"A": 18, "B": 15, "C": 15, "D": 15, "E": 4, "F": 12, "G": 12, "H": 12, "I": 12, "J": 12, "K": 12})

    # ------------------------------------------------------------ Par type d'asset
    ws = wb.create_sheet("Par type d'asset")
    _titre(ws, "Consommation par type d'asset", sous + " · les offres sont réparties selon leur composition")
    ty = logic.conso_par_type(lignes)
    _entetes(ws, 4, ["Type"] + entetes_mois + [f"Total {an}"], prev_from=1 + mr)
    ws.cell(row=4, column=14).fill = HEAD
    for i, (t, r) in enumerate(ty.iterrows()):
        row = 5 + i
        ws.cell(row=row, column=1, value=t).font = _f()
        for m in mois_cols:
            c = ws.cell(row=row, column=1 + m, value=float(r[m]))
            c.number_format = EUR0
        c = ws.cell(row=row, column=14, value=f"=SUM(B{row}:M{row})")
        c.number_format = EUR0
        c.font = _f(11, True)
    tr = 5 + len(ty)
    ws.cell(row=tr, column=1, value="TOTAL").font = _f(11, True)
    for col in range(1, 15):
        ws.cell(row=tr, column=col).fill = TOT
        if col >= 2:
            c = ws.cell(row=tr, column=col, value=f"=SUM({L(col)}5:{L(col)}{tr - 1})")
            c.number_format = EUR0
            c.font = _f(11, True)
    _largeurs(ws, {"A": 14, "N": 14})
    for k in range(12):
        ws.column_dimensions[L(2 + k)].width = 11

    # ------------------------------------------------------------ Détail (format long, pour tableaux croisés)
    ws = wb.create_sheet("Détail")
    _titre(ws, "Détail mensuel (une ligne par client, offre ou option et mois)",
           "Format adapté aux tableaux croisés dynamiques : Insertion > Tableau croisé dynamique.")
    hd = ["Mois", "ID client", "Client", "Entité", "Pays", "Élément", "Code", "Libellé", "Type", "Qté", "Montant", "Période"]
    _entetes(ws, 4, hd)
    row = 5
    for _, r in lignes.iterrows():
        for m in mois_cols:
            if r[m] == 0:
                continue
            vals = [datetime(an, m, 1), r.client_id, r.client, r.entite, r.pays, r.element, r.code, r.libelle,
                    r.type, float(r.qte or 0), float(r[m]), "Réel" if m <= mr else "Prévision"]
            for j, v in enumerate(vals):
                c = ws.cell(row=row, column=1 + j, value=v)
                if j == 0:
                    c.number_format = "mmm yyyy"
                if j == 10:
                    c.number_format = EUR
            row += 1
    _table(ws, "tRapportDetail", f"A4:L{max(row - 1, 5)}")
    _largeurs(ws, {"A": 11, "B": 9, "C": 26, "D": 15, "E": 11, "F": 9, "G": 8, "H": 30, "I": 10, "J": 6, "K": 13, "L": 10})
    ws.freeze_panes = "A5"

    # ------------------------------------------------------------ Relevés par client (option)
    if releves:
        for cid in conso.client_id:
            cl = clients.set_index("id").loc[cid]
            mine = lignes[lignes.client_id == cid]
            ws = wb.create_sheet(f"Relevé {cid}"[:31])
            _titre(ws, f"Relevé {an} — {cl.nom}", f"{cid} · {cl.entite} · {cl.pays} · responsable : {cl.responsable or '—'}")
            hd = ["Ligne", "Élément", "Code", "Libellé", "Qté", "Début", "Fin"] + entetes_mois + ["Total"]
            _entetes(ws, 4, hd, prev_from=7 + mr)
            ws.cell(row=4, column=20).fill = HEAD
            for i, (_, r) in enumerate(mine.iterrows()):
                row = 5 + i
                vals = [r.id, r.element, r.code, r.libelle, float(r.qte or 0), logic.d(r.debut), logic.d(r.fin)] + [float(r[m]) for m in mois_cols]
                for j, v in enumerate(vals):
                    c = ws.cell(row=row, column=1 + j, value=v)
                    if j in (5, 6):
                        c.number_format = "mmm yyyy"
                    if j >= 7:
                        c.number_format = EUR0
                c = ws.cell(row=row, column=20, value=f"=SUM(H{row}:S{row})")
                c.number_format = EUR0
                c.font = _f(11, True)
            tr = 5 + len(mine)
            ws.cell(row=tr, column=1, value="TOTAL").font = _f(11, True)
            for col in range(1, 21):
                ws.cell(row=tr, column=col).fill = TOT
                if col >= 8:
                    c = ws.cell(row=tr, column=col, value=f"=SUM({L(col)}5:{L(col)}{tr - 1})")
                    c.number_format = EUR0
                    c.font = _f(11, True)
            _largeurs(ws, {"A": 8, "B": 9, "C": 8, "D": 28, "E": 6, "F": 10, "G": 10, "T": 13})
            for k in range(12):
                ws.column_dimensions[L(8 + k)].width = 10

    for w in wb.worksheets:
        w.sheet_properties.tabColor = TEAL if not w.title.startswith("Relevé") else "9CC7BB"
        w.page_setup.orientation = "landscape"
        w.page_setup.fitToWidth = 1
        w.sheet_properties.pageSetUpPr.fitToPage = True
        w.page_setup.fitToHeight = 0
    wb.calculation.fullCalcOnLoad = True
    buf = BytesIO()
    wb.save(buf)
    return buf.getvalue()
